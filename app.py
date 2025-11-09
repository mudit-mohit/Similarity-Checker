import streamlit as st
import os
from docx import Document
from pdfminer.high_level import extract_text
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import base64
from io import BytesIO
from fpdf import FPDF
import unicodedata

st.set_page_config(page_title="Plagiarism Checker", page_icon="🔍", layout="wide")

st.title("🔍 Plagiarism Detection Tool")
st.markdown("Upload two documents to check for similarity")

def extract_text_from_file(uploaded_file):
    if uploaded_file.type == "application/pdf":
        return extract_text(uploaded_file)
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(uploaded_file)
        return '\n'.join([para.text for para in doc.paragraphs])
    else:
        raise ValueError("Unsupported file format")

def calculate_similarity(text1, text2):
    vectorizer = TfidfVectorizer()
    vectors = vectorizer.fit_transform([text1, text2])
    return cosine_similarity(vectors[0:1], vectors[1:2])[0][0]

def create_pdf_report(path1, path2, score, text1, text2):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    def safe_text(text):
        if not text:
            return ""
        text = str(text)
        text = unicodedata.normalize('NFKD', text)
        text = text.encode('latin-1', 'replace').decode('latin-1')
        return text

    pdf.cell(200, 10, txt="Plagiarism Detection Report", ln=True, align='C')
    pdf.ln(10)
    pdf.cell(200, 10, txt=f"Document 1: {safe_text(path1)}", ln=True)
    pdf.cell(200, 10, txt=f"Document 2: {safe_text(path2)}", ln=True)
    pdf.cell(200, 10, txt=f"Similarity Score: {score}%", ln=True)
    pdf.ln(10)
    pdf.multi_cell(0, 10, txt="--- Document 1 Content ---")
    safe_text1 = safe_text(text1[:1000] + "..." if len(text1) > 1000 else text1)
    pdf.multi_cell(0, 10, txt=safe_text1)
    pdf.ln(5)
    pdf.multi_cell(0, 10, txt="--- Document 2 Content ---")
    safe_text2 = safe_text(text2[:1000] + "..." if len(text2) > 1000 else text2)
    pdf.multi_cell(0, 10, txt=safe_text2)

    return pdf.output(dest='S').encode('latin-1')

# File upload section
col1, col2 = st.columns(2)

with col1:
    st.subheader("Document 1")
    doc1 = st.file_uploader("Upload first document", type=['pdf', 'docx'], key="doc1")

with col2:
    st.subheader("Document 2")
    doc2 = st.file_uploader("Upload second document", type=['pdf', 'docx'], key="doc2")

# Check plagiarism button
if st.button("✅ Check Plagiarism", type="primary"):
    if doc1 and doc2:
        with st.spinner("Analyzing documents..."):
            try:
                text1 = extract_text_from_file(doc1)
                text2 = extract_text_from_file(doc2)
                
                if not text1.strip() or not text2.strip():
                    st.error("One or both documents are empty.")
                else:
                    score = calculate_similarity(text1, text2)
                    plagiarism_score = round(score * 100, 2)
                    
                    # Display result
                    st.success(f"🔍 Plagiarism Score: **{plagiarism_score}%**")
                    
                    # Export option
                    pdf_data = create_pdf_report(doc1.name, doc2.name, plagiarism_score, text1, text2)
                    st.download_button(
                        label="📥 Download PDF Report",
                        data=pdf_data,
                        file_name="plagiarism_report.pdf",
                        mime="application/pdf"
                    )
                    
                    # Show text preview
                    with st.expander("View Document Contents"):
                        tab1, tab2 = st.tabs(["Document 1", "Document 2"])
                        with tab1:
                            st.text_area("Document 1 Content", text1[:2000] + "..." if len(text1) > 2000 else text1, height=200)
                        with tab2:
                            st.text_area("Document 2 Content", text2[:2000] + "..." if len(text2) > 2000 else text2, height=200)
                            
            except Exception as e:
                st.error(f"Error: {str(e)}")
    else:
        st.warning("Please upload both documents first.")

st.markdown("---")
st.markdown("Built with Streamlit | AI/ML Plagiarism Detection")